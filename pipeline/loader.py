import logging
from pathlib import Path
from typing import List, Dict, Any
from pypdf import PdfReader
import openpyxl

# Silence pypdf font warnings
logging.getLogger("pypdf").setLevel(logging.ERROR)


class DocumentLoader:
    """Extracts text and metadata from PDFs, Excel sheets, and text documents."""

    @staticmethod
    def load_pdf(file_path: Path) -> List[Dict[str, Any]]:
        """Extract text page by page from a PDF."""
        docs = []
        try:
            reader = PdfReader(str(file_path))
            for page_idx, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    docs.append({
                        "content": text.strip(),
                        "metadata": {
                            "source": file_path.name,
                            "file_path": str(file_path),
                            "file_type": "pdf",
                            "page": page_idx + 1,
                            "total_pages": len(reader.pages),
                        }
                    })
        except Exception as e:
            print(f"Error reading PDF {file_path.name}: {e}")
        return docs

    @staticmethod
    def load_excel(file_path: Path) -> List[Dict[str, Any]]:
        """Extract text from Excel worksheets."""
        docs = []
        try:
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue

                headers = [str(h).strip() if h is not None else f"col_{idx}" for idx, h in enumerate(rows[0])]
                sheet_lines = []

                for row_idx, row in enumerate(rows[1:], start=2):
                    row_items = []
                    for h, val in zip(headers, row):
                        if val is not None and str(val).strip():
                            row_items.append(f"{h}: {str(val).strip()}")
                    if row_items:
                        sheet_lines.append(f"Row {row_idx}: " + ", ".join(row_items))

                if sheet_lines:
                    text_content = f"Sheet: {sheet_name}\n" + "\n".join(sheet_lines)
                    docs.append({
                        "content": text_content,
                        "metadata": {
                            "source": file_path.name,
                            "file_path": str(file_path),
                            "file_type": "xlsx",
                            "sheet": sheet_name,
                            "total_rows": len(rows),
                        }
                    })
        except Exception as e:
            print(f"Error reading Excel {file_path.name}: {e}")
        return docs

    @classmethod
    def load_directory(cls, dir_path: Path) -> List[Dict[str, Any]]:
        """Loads all supported documents from the target directory."""
        if not dir_path.exists():
            raise FileNotFoundError(f"Directory not found: {dir_path}")

        all_docs = []
        supported_files = sorted(list(dir_path.glob("*.*")))

        for file_path in supported_files:
            suffix = file_path.suffix.lower()
            if suffix == ".pdf":
                docs = cls.load_pdf(file_path)
                all_docs.extend(docs)
                print(f"Loaded {len(docs)} pages from {file_path.name}")
            elif suffix in [".xlsx", ".xls"]:
                docs = cls.load_excel(file_path)
                all_docs.extend(docs)
                print(f"Loaded {len(docs)} sheet sections from {file_path.name}")
            elif suffix in [".txt", ".md"]:
                try:
                    text = file_path.read_text(encoding="utf-8")
                    if text.strip():
                        all_docs.append({
                            "content": text.strip(),
                            "metadata": {
                                "source": file_path.name,
                                "file_path": str(file_path),
                                "file_type": suffix[1:],
                            }
                        })
                        print(f"Loaded text file {file_path.name}")
                except Exception as e:
                    print(f"Error reading text file {file_path.name}: {e}")

        return all_docs
